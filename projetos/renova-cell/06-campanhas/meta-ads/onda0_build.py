"""Gera a planilha de montagem da Onda 0 (onda-0-montagem.xlsx), os anúncios em texto
(onda-0-anuncios.md) e a seção 8.3 de onda-0-estrutura.md a partir de uma única fonte de dados.

Rodar: python3 onda0_build.py   (na pasta 06-campanhas/meta-ads)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
DOMAIN = 'https://vitalrise.vercel.app'
URL_PARAMS = 'utm_source=meta&utm_medium=paid&utm_campaign={{campaign.name}}&utm_term={{adset.name}}&utm_content={{ad.name}}'
FOOTER = ('Individual results vary. These statements have not been evaluated by the Food and Drug Administration. '
          'This product is not intended to diagnose, treat, cure, or prevent any disease.')

# ---------------------------------------------------------------- campanhas
CAMPAIGNS = [
    {
        'name': 'rc_us_c0_aquecimento',
        'when': 'D0 a D3',
        'objective': 'Engajamento → Visualizações de vídeo (ThruPlay)',
        'budget': 'Orçamento do conjunto (ABO)',
        'daily': 15,
        'days': 4,
        'event': '— (ThruPlay)',
        'why': 'Conta e pixel novos: entrega estável, primeiros públicos de vídeo e engajamento acumulado nas mesmas publicações que viram teste.',
    },
    {
        'name': 'rc_us_c1_teste',
        'when': 'D4 a D17',
        'objective': 'Vendas → site',
        'budget': 'Orçamento do conjunto (ABO)',
        'daily': 100,
        'days': 14,
        'event': 'Lead (quiz concluído). Caminho B: Visualizações da página de destino',
        'why': 'Comparar 5 conceitos com a mesma verba. Lead dá sinal rápido; a compra é julgada pela Stripe (utm_content).',
    },
    {
        'name': 'rc_us_c2_escala',
        'when': 'D11 a D21',
        'objective': 'Vendas → site (campanha Advantage+ de vendas)',
        'budget': 'Orçamento da campanha',
        'daily': 60,
        'days': 11,
        'event': 'Purchase se houver ≥ 10 compras na semana; senão InitiateCheckout. Caminho B: Visualizações da página',
        'why': 'Só os vencedores do teste (mesmas publicações), um conjunto amplo.',
    },
    {
        'name': 'rc_us_c3_remarketing',
        'when': 'D4 a D21',
        'objective': 'Vendas → site',
        'budget': 'Orçamento do conjunto (ABO)',
        'daily': 15,
        'days': 18,
        'event': 'Purchase (caminho A) / Visualizações da página (caminho B)',
        'why': 'Quem começou o quiz, viu a oferta ou iniciou o checkout e não comprou.',
    },
]

# ---------------------------------------------------------------- conjuntos
BASE_TARGET = 'EUA · 35–65+ · todos os gêneros · inglês (todos os idiomas se o público ficar pequeno) · Público Advantage+ com 35+ como sugestão'
PLACEMENTS = 'Posicionamentos Advantage+ (Instagram + Facebook), excluindo Audience Network e Messenger'
ADSETS = [
    ('rc_us_c0_aquecimento', 'aq_video_amplo', BASE_TARGET, PLACEMENTS, 15, 'ThruPlay', 'ig01a, ig06a, ig04a (mesmas publicações do teste)'),
    ('rc_us_c1_teste', 'ig01_ritual', BASE_TARGET, PLACEMENTS, 20, 'Lead', 'ig01a, ig01b, ig01c'),
    ('rc_us_c1_teste', 'ig03_unboxing', BASE_TARGET, PLACEMENTS, 20, 'Lead', 'ig03a, ig03b, ig03c'),
    ('rc_us_c1_teste', 'ig04_quiz', BASE_TARGET, PLACEMENTS, 20, 'Lead', 'ig04a, ig04b, ig04c'),
    ('rc_us_c1_teste', 'ig06_a_dois', BASE_TARGET, PLACEMENTS, 20, 'Lead', 'ig06a, ig06b, ig06c'),
    ('rc_us_c1_teste', 'ig10_criador (reserva: ig05_assinatura)', BASE_TARGET, PLACEMENTS, 20, 'Lead',
     'ig10a, ig10b, ig10c (anúncio de parceria). Sem criador pronto até D4: ig05a, ig05b, ig05c'),
    ('rc_us_c2_escala', 'escala_amplo', 'EUA · 18–65+ (Advantage+ decide) · exclui compradores 180 d', PLACEMENTS, 60,
     'Purchase / InitiateCheckout', 'Os 3–6 vencedores do C1 (usar a publicação existente: Post ID)'),
    ('rc_us_c3_remarketing', 'rt1_quiz_oferta_14d',
     'Público: iniciou o quiz 14 d + viu a oferta 14 d + iniciou checkout 7 d · exclui compradores 180 d', PLACEMENTS, 10,
     'Purchase', 'ig08a, ig12a, ig03b, ig13a'),
    ('rc_us_c3_remarketing', 'rt2_engajou_30d',
     'Público: engajou com o Instagram 30 d + viu 50% de vídeo 30 d · exclui quem iniciou o quiz 30 d e compradores', PLACEMENTS, 5,
     'Purchase', 'ig04a, ig08a'),
]

# ---------------------------------------------------------------- anúncios
QUIZ = f'{DOMAIN}/quiz'
OFFER = f'{DOMAIN}/offer'
CONCEPTS = {
    'ig01': {
        'title': 'Ritual da manhã',
        'format': 'Vídeo 9:16 (Reels/Stories) + corte 4:5 (Feed) · 15–20 s',
        'scenes': 'Colher no pó (câmera lenta, som de ASMR) → mistura no café → gole na janela → produto na bancada → cartão final "Free shipping · 2 business days · Take the 60-sec quiz". Selo "Dramatization" se for ator.',
        'primary': ('A 30-second daily ritual: one serving of collagen + peptides in coffee, water or a smoothie. No hormones, no stimulants.\n\n'
                    'Free shipping, delivered in up to 2 business days in plain, discreet packaging.\n\n'
                    'Take the 60-second quiz to get the plan that matches your answers. Buy once or subscribe, cancel anytime. 30-day money-back guarantee.'),
        'pt': 'Um ritual diário de 30 segundos: uma dose de colágeno + peptídeos no café, na água ou na vitamina. Sem hormônios, sem estimulantes. Frete grátis, entrega em até 2 dias úteis em embalagem discreta. Faça o quiz de 60 segundos. Compra única ou assinatura, cancele quando quiser. Garantia de 30 dias.',
        'hooks': [('One scoop. Every morning. That\'s it.', 'Your 30-second daily ritual'),
                  ('A 30-second morning ritual', 'Free shipping · 2 business days'),
                  ('The supplement that fits in a coffee mug', 'Find your plan in 60 seconds')],
        'description': 'Buy once or subscribe · cancel anytime',
        'cta': 'Saiba mais (Learn more)',
        'url': QUIZ,
    },
    'ig03': {
        'title': 'Unboxing discreto + 2 dias',
        'format': 'Vídeo 9:16 só com mãos · 12–15 s',
        'scenes': 'Caixa neutra na porta → close da caixa sem marca → abre → produto → mistura na água → cartão final. Gancho C só depois de comprovar a entrega real.',
        'primary': ('Some things are nobody else\'s business. VitalRise ships in a plain box with no product name on the outside.\n\n'
                    'Free shipping, delivered in up to 2 business days with tracking. Buy once, or subscribe and it arrives again when you run out. Cancel anytime.\n\n'
                    'Take the 60-second quiz to find your plan. 30-day money-back guarantee.'),
        'pt': 'Tem coisa que não é da conta de ninguém. O VitalRise chega numa caixa neutra, sem o nome do produto por fora. Frete grátis, entrega em até 2 dias úteis com rastreio. Compra única ou assinatura que chega de novo quando acaba. Faça o quiz de 60 segundos. Garantia de 30 dias.',
        'hooks': [('No logo on the box. On purpose.', 'Free shipping · discreet box'),
                  ('Free shipping. 2 business days. Plain box.', 'Delivered in up to 2 business days'),
                  ('Ordered Monday. Here Wednesday.', 'Plain box. Fast delivery.')],
        'description': 'Tracked · discreet · free',
        'cta': 'Saiba mais (Learn more)',
        'url': QUIZ,
    },
    'ig04': {
        'title': 'O quiz é o anúncio',
        'format': 'Vídeo 9:16: gravação de tela do quiz real (rápida, com cliques) · 12–15 s',
        'scenes': 'Dedo respondendo as 8 perguntas em velocidade 2× → tela "Analysis complete" → plano recomendado → cartão final.',
        'primary': ('8 quick questions, about 60 seconds, and you get a daily collagen + peptide plan matched to your answers: 1 month, 3 months or 6 months.\n\n'
                    'Free shipping in up to 2 business days, discreet packaging, 30-day money-back guarantee.'),
        'pt': '8 perguntas rápidas, cerca de 60 segundos, e você recebe um plano diário de colágeno + peptídeos de acordo com as suas respostas: 1, 3 ou 6 meses. Frete grátis em até 2 dias úteis, embalagem discreta, garantia de 30 dias.',
        'hooks': [('Answer 8 questions. Get your collagen plan.', 'Take the 60-second quiz'),
                  ('Which plan fits your routine? 60 seconds.', 'Get your plan in 60 seconds'),
                  ('Coffee, water or smoothie? Start here.', '8 questions. 1 plan.')],
        'description': 'Free shipping · 2 business days',
        'cta': 'Saiba mais (Learn more)',
        'url': QUIZ,
    },
    'ig06': {
        'title': 'Ritual a dois',
        'format': 'Vídeo 9:16 + corte 4:5 · 15–20 s',
        'scenes': 'Casal 40–55 na cozinha de manhã, duas canecas, risada, abraço rápido; nada de cama nem toque sensual. Selo "Dramatization" se forem atores.',
        'primary': ('A daily ritual for two: one serving of collagen + peptides each, in coffee or a smoothie. No hormones, no stimulants.\n\n'
                    'Free shipping, delivered in up to 2 business days in discreet packaging. Buy once or subscribe, cancel anytime.\n\n'
                    'Take the 60-second quiz to find your plan.'),
        'pt': 'Um ritual diário a dois: uma dose de colágeno + peptídeos para cada um, no café ou na vitamina. Sem hormônios, sem estimulantes. Frete grátis em até 2 dias úteis, embalagem discreta. Compra única ou assinatura. Faça o quiz de 60 segundos.',
        'hooks': [('Our morning thing.', 'A daily ritual for two'),
                  ('Two mugs. One ritual.', 'Free shipping · 2 business days'),
                  ('Self-care is better together.', 'Take the 60-second quiz')],
        'description': 'No hormones · no stimulants',
        'cta': 'Saiba mais (Learn more)',
        'url': QUIZ,
    },
    'ig10': {
        'title': 'Criador (anúncio de parceria)',
        'format': 'Reels do criador, 9:16, 20–40 s, publicado no perfil dele e impulsionado como anúncio de parceria',
        'scenes': 'Rotina real do criador + caixa discreta chegando + preparo no café. Três aberturas gravadas (A/B/C). Rótulo "Paid partnership" + "#ad" + divulgação falada nos primeiros segundos.',
        'primary': ('[Legenda do próprio criador, aprovada pela marca — só experiência real, sem resultado.] Example: "I\'ve been adding VitalRise to my morning coffee for 4 weeks. One scoop and done. Free shipping, and it came in a plain box. #ad"'),
        'pt': '[Legenda do criador.] Exemplo: "Faz 4 semanas que coloco VitalRise no meu café da manhã. Uma dose e pronto. Frete grátis, e chegou numa caixa neutra. #publi" — só se for verdade.',
        'hooks': [('[Abertura A do criador]', 'Take the 60-second quiz'),
                  ('[Abertura B do criador]', 'Free shipping · 2 business days'),
                  ('[Abertura C do criador]', 'Find your plan in 60 seconds')],
        'description': 'Paid partnership',
        'cta': 'Saiba mais (Learn more)',
        'url': QUIZ,
    },
    'ig05': {
        'title': 'Chega quando acaba (reserva do ig10)',
        'format': 'Vídeo 9:16 animado (calendário + produto) · 15 s',
        'scenes': 'Calendário → pacote chegando no fim do estoque (1, 3 ou 6 meses) → "Cancel anytime from your receipt email" → cartão final.',
        'primary': ('Choose your plan: 1 month, 3 months or 6 months. Buy once, or subscribe and it arrives again right when you run out. Same price, cancel anytime from your receipt email.\n\n'
                    'Free shipping in up to 2 business days, discreet packaging. 30-day money-back guarantee.'),
        'pt': 'Escolha 1, 3 ou 6 meses. Compre uma vez ou assine e ele chega de novo quando acabar. Mesmo preço, cancele quando quiser pelo e-mail do recibo. Frete grátis em até 2 dias úteis, embalagem discreta. Garantia de 30 dias.',
        'hooks': [('It shows up right when you run out.', 'Subscribe or buy once'),
                  ('1, 3 or 6 months. You choose.', 'Same price · cancel anytime'),
                  ('Subscribe or buy once. Same price.', 'Free shipping · 2 business days')],
        'description': 'Cancel anytime',
        'cta': 'Saiba mais (Learn more)',
        'url': QUIZ,
    },
}
RT_ADS = [
    {
        'name': 'ig08a', 'title': 'Garantia de 30 dias', 'format': 'Imagem 4:5 + versão 9:16 (Stories)',
        'onscreen': 'Arte: produto + "30-day money-back guarantee · Free shipping · 2 business days"',
        'primary': ('Still thinking about it? Try VitalRise for 30 days. If it\'s not for you, email us and we\'ll refund your first order.\n\n'
                    'Free shipping, delivered in up to 2 business days in discreet packaging.'),
        'headline': '30-day money-back guarantee', 'description': 'Free shipping · discreet box', 'cta': 'Comprar agora (Shop now)', 'url': OFFER,
    },
    {
        'name': 'ig12a', 'title': 'O que vem na caixa', 'format': 'Imagem 4:5 (flat lay) + 9:16',
        'onscreen': 'Arte: "What\'s in the box" · "Your plan (1, 3 or 6 months)" · "Plain outer box" · "Free shipping" · "30-day guarantee"',
        'primary': ('Your plan, a plain outer box and free delivery in up to 2 business days. Buy once or subscribe, cancel anytime. 30-day money-back guarantee.'),
        'headline': 'Free shipping · 2 business days', 'description': 'Buy once or subscribe', 'cta': 'Comprar agora (Shop now)', 'url': OFFER,
    },
    {
        'name': 'ig03b', 'title': 'Unboxing (reaproveitado do teste)', 'format': 'Mesma publicação do C1 (Post ID)',
        'onscreen': 'Gancho B do ig03', 'primary': '(mesmo texto do ig03)', 'headline': 'Delivered in up to 2 business days',
        'description': 'Tracked · discreet · free', 'cta': 'Comprar agora (Shop now)', 'url': OFFER,
    },
    {
        'name': 'ig13a', 'title': 'Cupom VITAL10 (quem iniciou o checkout)', 'format': 'Imagem 4:5 + 9:16',
        'onscreen': 'Arte: "10% off your first payment · code VITAL10" + produto',
        'primary': ('Your plan is still waiting. Use code VITAL10 for 10% off your first payment.\n\n'
                    'Free shipping in up to 2 business days, discreet packaging, 30-day money-back guarantee.'),
        'headline': '10% off with VITAL10', 'description': 'First payment only', 'cta': 'Obter oferta (Get offer)', 'url': OFFER,
    },
]

TEST_CONCEPTS = ['ig01', 'ig03', 'ig04', 'ig06', 'ig10', 'ig05']


def ad_rows():
    rows = []
    for key in TEST_CONCEPTS:
        c = CONCEPTS[key]
        for i, (hook, headline) in enumerate(c['hooks']):
            name = f'{key}{"abc"[i]}'
            note = ''
            if key == 'ig05':
                note = 'Reserva: só sobe se o ig10 (criador) não estiver pronto até D4.'
            if key == 'ig03' and i == 2:
                note = 'Só depois de comprovar a entrega real (D-7 a D-3). Antes disso, use "Tracked. Fast. Discreet."'
            rows.append({
                'campaign': 'rc_us_c1_teste', 'adset': f'{key}_' + {'ig01': 'ritual', 'ig03': 'unboxing', 'ig04': 'quiz', 'ig06': 'a_dois', 'ig10': 'criador', 'ig05': 'assinatura'}[key],
                'name': name, 'concept': c['title'], 'format': c['format'], 'onscreen': hook,
                'primary': c['primary'] + '\n\n' + FOOTER, 'headline': headline, 'description': c['description'],
                'cta': c['cta'], 'url': c['url'], 'note': note,
            })
    for ad in RT_ADS:
        adset = 'rt1_quiz_oferta_14d'
        rows.append({
            'campaign': 'rc_us_c3_remarketing', 'adset': adset, 'name': ad['name'], 'concept': ad['title'], 'format': ad['format'],
            'onscreen': ad['onscreen'], 'primary': (ad['primary'] if ad['primary'].startswith('(') else ad['primary'] + '\n\n' + FOOTER),
            'headline': ad['headline'], 'description': ad['description'], 'cta': ad['cta'], 'url': ad['url'],
            'note': 'Também no rt2_engajou_30d.' if ad['name'] == 'ig08a' else ('Só para quem iniciou o checkout: use um conjunto à parte se o público passar de ~1.000 pessoas.' if ad['name'] == 'ig13a' else ''),
        })
    return rows


AUDIENCES = [
    ('aud_quiz_start_14d', 'Site', 'Evento quiz_start (ou URL contém /quiz)', '14 dias', 'Remarketing RT1'),
    ('aud_quiz_start_30d', 'Site', 'Evento quiz_start', '30 dias', 'Exclusão do RT2'),
    ('aud_lead_30d', 'Site', 'Evento Lead (quiz concluído)', '30 dias', 'Leitura; futuro lookalike'),
    ('aud_offer_14d', 'Site', 'URL contém /offer', '14 dias', 'Remarketing RT1'),
    ('aud_checkout_7d', 'Site', 'Evento InitiateCheckout', '7 dias', 'Remarketing RT1 / ig13a'),
    ('aud_compradores_180d', 'Site + lista', 'Evento Purchase + lista de clientes da Stripe (e-mail)', '180 dias', 'Exclusão em todas as campanhas'),
    ('aud_ig_engajou_30d', 'Perfil do Instagram', 'Todas as pessoas que interagiram com a conta', '30 dias', 'Remarketing RT2'),
    ('aud_video50_30d', 'Vídeo', 'Viram 50% de qualquer vídeo do C0/C1', '30 dias', 'Remarketing RT2'),
    ('lal_compradores_1pct', 'Semelhante', 'aud_compradores_180d, EUA, 1%', '—', 'Só com ≥ 100 compradores (provavelmente na Onda 1)'),
]

CHECKLIST = [
    ('Site', 'SHIP_COUNTRIES=US e FAST_DELIVERY_COUNTRIES=US no ambiente de produção', 'D-10'),
    ('Site', 'Chaves Stripe live + npm run stripe:setup (preços e cupom VITAL10)', 'D-10'),
    ('Site', 'NEXT_PUBLIC_META_PIXEL_ID + META_CAPI_TOKEN; supabase/schema.sql rodado de novo', 'D-10'),
    ('Conta', 'Portfólio empresarial verificado, 2 administradores com 2FA', 'D-10'),
    ('Conta', 'Conta de anúncios USD, fuso de quem opera, forma de pagamento + reserva', 'D-10'),
    ('Conta', 'Página do Facebook + Instagram profissional conectados; anúncios de parceria ativados no Instagram', 'D-10'),
    ('Conta', 'Domínio verificado (Segurança da marca → Domínios)', 'D-9'),
    ('Conta', 'Verificar se o conjunto de dados (pixel) foi classificado como saúde e bem-estar → caminho A ou B', 'D-9'),
    ('Rastreio', 'Eventos de teste: PageView, quiz_start, Lead, ViewContent, InitiateCheckout, Purchase (pixel + servidor deduplicados)', 'D-7'),
    ('Rastreio', 'Compra real de teste (reembolsar depois): Purchase aparece 1 vez, com valor e moeda', 'D-7'),
    ('Rastreio', 'Remover META_TEST_EVENT_CODE depois dos testes', 'D-6'),
    ('Logística', '10 pedidos reais para endereços em estados diferentes; anotar dias úteis até a entrega', 'D-7 a D-3'),
    ('Criativos', '15 vídeos do teste (5 conceitos × 3 ganchos) + 4 peças de remarketing aprovados', 'D-3'),
    ('Criativos', 'Publicar o post orgânico de cada anúncio e anotar o Post ID (para reaproveitar o engajamento)', 'D-2'),
    ('Instagram', 'Grade inicial de 9 posts + destaques + bio com link do quiz', 'D-2'),
    ('Públicos', 'Criar os 8 públicos da aba Públicos', 'D-1'),
    ('Campanhas', 'Subir C0, C1 e C3 PAUSADAS, revisar políticas, conferir URLs (abrir cada uma)', 'D-1'),
    ('Campanhas', 'Ligar C0', 'D0'),
    ('Campanhas', 'Ligar C1 e C3', 'D4'),
    ('Campanhas', 'Montar C2 com os vencedores', 'D11'),
]


def build_xlsx():
    wb = Workbook()
    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='0E5C4A')
    wrap = Alignment(wrap_text=True, vertical='top')

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = head, fill, wrap
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        ws.freeze_panes = 'A2'
        return ws

    wb.remove(wb.active)
    ws = sheet('Campanhas', ['Campanha', 'Quando', 'Objetivo', 'Tipo de orçamento', 'US$/dia', 'Dias', 'Total US$', 'Evento de otimização', 'Para quê'],
               [(c['name'], c['when'], c['objective'], c['budget'], c['daily'], c['days'], None, c['event'], c['why']) for c in CAMPAIGNS],
               [26, 12, 30, 22, 9, 7, 10, 40, 60])
    for r in range(2, 2 + len(CAMPAIGNS)):
        ws[f'G{r}'] = f'=E{r}*F{r}'
    total = 2 + len(CAMPAIGNS)
    ws[f'A{total}'] = 'Total da Onda 0 (exemplo)'
    ws[f'A{total}'].font = Font(bold=True)
    ws[f'G{total}'] = f'=SUM(G2:G{total - 1})'
    ws[f'G{total}'].font = Font(bold=True)
    ws[f'A{total + 1}'] = 'Valores de exemplo. Ajuste US$/dia ao que você aceita investir; a regra de cada conjunto de teste é verba diária ≥ CPA alvo provisório.'
    for r in range(2, total + 1):
        ws[f'E{r}'].number_format = ws[f'G{r}'].number_format = '"US$" #,##0'

    sheet('Conjuntos', ['Campanha', 'Conjunto', 'Público', 'Posicionamentos', 'US$/dia', 'Otimização', 'Anúncios'],
          ADSETS, [24, 22, 50, 40, 9, 22, 50])
    sheet('Anúncios', ['Campanha', 'Conjunto', 'Anúncio', 'Conceito', 'Formato', 'Texto na tela (gancho, 0–2 s)', 'Texto principal (EN)',
                       'Título', 'Descrição', 'CTA', 'URL do site', 'Parâmetros de URL', 'Observação'],
          [(a['campaign'], a['adset'], a['name'], a['concept'], a['format'], a['onscreen'], a['primary'], a['headline'], a['description'],
            a['cta'], a['url'], URL_PARAMS, a['note']) for a in ad_rows()],
          [20, 20, 9, 20, 28, 34, 70, 30, 26, 22, 26, 40, 36])
    sheet('Públicos', ['Nome', 'Origem', 'Regra', 'Janela', 'Uso'], AUDIENCES, [24, 18, 50, 10, 40])
    sheet('Checklist', ['Área', 'Item', 'Prazo', 'Feito?'], [(a, b, c, '') for a, b, c in CHECKLIST], [12, 80, 12, 10])
    wb.save(HERE / 'onda-0-montagem.xlsx')


def build_md():
    out = []
    for key in TEST_CONCEPTS:
        c = CONCEPTS[key]
        adset = next(a['adset'] for a in ad_rows() if a['name'].startswith(key))
        out.append(f'#### {key.upper()} — {c["title"]} · conjunto `{adset}`\n')
        out.append(f'- **Formato:** {c["format"]}')
        out.append(f'- **Cenas:** {c["scenes"]}')
        out.append(f'- **Texto principal (EN, igual nos 3 anúncios):**\n')
        out.append('  > ' + c['primary'].replace('\n\n', '\n  >\n  > ') + '\n  >\n  > ' + FOOTER + '\n')
        out.append(f'  *{c["pt"]}*\n')
        out.append('| Anúncio | Gancho na tela (0–2 s) | Título |')
        out.append('|---|---|---|')
        for i, (hook, headline) in enumerate(c['hooks']):
            out.append(f'| `{key}{"abc"[i]}` | "{hook}" | "{headline}" |')
        out.append(f'\n- **Descrição:** "{c["description"]}" · **CTA:** {c["cta"]} · **URL:** `{c["url"]}`\n')
    out.append('#### Remarketing · conjuntos `rt1_quiz_oferta_14d` e `rt2_engajou_30d`\n')
    out.append('| Anúncio | Peça | Formato | Texto na arte | Texto principal (EN) | Título | CTA | URL |')
    out.append('|---|---|---|---|---|---|---|---|')
    for ad in RT_ADS:
        primary = ad['primary'].replace('\n\n', ' ')
        out.append(f'| `{ad["name"]}` | {ad["title"]} | {ad["format"]} | {ad["onscreen"]} | {primary} | "{ad["headline"]}" | {ad["cta"]} | `{ad["url"]}` |')
    out.append('\nTodo texto principal de remarketing termina com o mesmo rodapé da FDA.')
    text = '\n'.join(out) + '\n'
    (HERE / 'onda-0-anuncios.md').write_text(text)
    # Same section inside the main document, between its markers.
    doc = HERE / 'onda-0-estrutura.md'
    start, end = '<!-- ANUNCIOS:INICIO', '<!-- ANUNCIOS:FIM -->'
    s = doc.read_text()
    i = s.index('\n', s.index(start)) + 1
    j = s.index(end)
    doc.write_text(s[:i] + text + s[j:])


if __name__ == '__main__':
    build_xlsx()
    build_md()
    print('ok:', len(ad_rows()), 'anúncios')
